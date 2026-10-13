#!/usr/bin/env python3
"""
transcriber.py — lokalna transkrypcja nagrań (głównie wykładów PL). Wszystko w Pythonie, bez kluczy API.

  python transcriber.py wyklad.mp4              # transkrypcja z timestampami
  python transcriber.py wyklad.mp4 --speakers   # + podział na mówców
  python transcriber.py wyklad.mp4 --extract    # + slajdy (istotne zmiany obrazu) jako JPG
  python transcriber.py wyklad.mp4 --summary    # + opisy slajdów (Qwen3-VL) + notatka .md
  python transcriber.py wyklad.mp4 --web        # + notatka uzupełniona wiedzą z internetu
  python transcriber.py wyklad.mp4 --materials slajdy.pdf notatki.docx   # + czyste dane z prezentacji
  python transcriber.py wyklad.mp4 --all        # wszystko naraz
  python transcriber.py wyklad_wyniki/wyklad_notatka.md   # tylko konwersja gotowej notatki do PDF

Instalacja (jednorazowo):
  python3 -m venv .venv && source .venv/bin/activate     # Windows: .venv\\Scripts\\activate
  pip install -r requirements.txt

Modele pobierają się przy pierwszym użyciu do ./models obok skryptu — tylko te potrzebne:
  zawsze        Whisper (transkrypcja)
  --speakers    pyannote community-1
  --summary     Qwen3-VL (opis slajdów + notatka): MLX na Macu, 4-bit na NVIDIA, bf16 na CPU
"""
from __future__ import annotations

import argparse
import gc
import json
import os
import platform
import re
import sys
import tempfile
import threading
import time
from contextlib import contextmanager
from dataclasses import asdict, dataclass
from pathlib import Path

import numpy as np
from tqdm import tqdm

# ============================================================================
# Konfiguracja
# ============================================================================

SCRIPT_DIR = Path(__file__).resolve().parent
MODELS_DIR = SCRIPT_DIR / "models"

SAMPLE_RATE = 16000
VIDEO_EXTS = {".mp4", ".mkv", ".mov", ".avi", ".webm", ".m4v", ".flv", ".wmv"}
AUDIO_EXTS = {".mp3", ".wav", ".m4a", ".flac", ".ogg", ".opus", ".aac"}

DEFAULT_WHISPER = "large-v3-turbo"
# Model mówców: pyannote community-1 (CC-BY-4.0, autor: pyannote / Hervé Bredin).
# Pobieramy publiczną kopię (bez tokenu), przypiętą do konkretnego commita i sprawdzaną sumami SHA-256 –
# pliki są identyczne bajt w bajt z oficjalnym repozytorium pyannote/speaker-diarization-community-1.
PYANNOTE_REPO = "pyannote/speaker-diarization-community-1"  # oficjalne (wymaga tokenu) – zapasowo
PYANNOTE_MIRROR = "pyannote-community/speaker-diarization-community-1"
PYANNOTE_REVISION = "8a527374977391da736e0daaef26855d949d9685"
PYANNOTE_SHA256 = {
    "embedding/pytorch_model.bin": "6f10ff60898a1d185fa22e1d11e0bfa8a92efec811f11bca48cb8cafebefd929",
    "segmentation/pytorch_model.bin": "7ad24338d844fb95985486eb1a464e32d229f6d7a03c9abe60f978bacf3f816e",
    "plda/plda.npz": "9b77bcd840692710dd3496f62ecfeed8d8e5f002fd991b785079b244eab7d255",
    "plda/xvec_transform.npz": "325f1ce8e48f7e55e9c8aa47e05d2766b7c48c4b25b8de8dd751e7a4cc5fbe8f",
}

# nazwy modeli faster-whisper -> repozytoria MLX (Apple Silicon)
MLX_WHISPER = {
    "large-v3-turbo": "mlx-community/whisper-large-v3-turbo",
    "turbo": "mlx-community/whisper-large-v3-turbo",
    "large-v3": "mlx-community/whisper-large-v3-mlx",
    "large-v2": "mlx-community/whisper-large-v2-mlx",
    "medium": "mlx-community/whisper-medium-mlx",
    "small": "mlx-community/whisper-small-mlx",
    "base": "mlx-community/whisper-base-mlx",
    "tiny": "mlx-community/whisper-tiny-mlx",
}

# Qwen3-VL: wariant dobrany do sprzętu (backend -> rozmiar -> repozytorium HF)
QWEN = {
    "mlx": {  # Apple Silicon, 4-bit MLX
        "2b": "mlx-community/Qwen3-VL-2B-Instruct-4bit",
        "4b": "mlx-community/Qwen3-VL-4B-Instruct-4bit",
        "8b": "mlx-community/Qwen3-VL-8B-Instruct-4bit",
    },
    "cuda": {  # NVIDIA, 4-bit bitsandbytes (prekwantyzowane – mniejszy download)
        "2b": "unsloth/Qwen3-VL-2B-Instruct-unsloth-bnb-4bit",
        "4b": "unsloth/Qwen3-VL-4B-Instruct-unsloth-bnb-4bit",
        "8b": "unsloth/Qwen3-VL-8B-Instruct-unsloth-bnb-4bit",
    },
    "cpu": {  # bez akceleracji, bf16
        "2b": "Qwen/Qwen3-VL-2B-Instruct",
        "4b": "Qwen/Qwen3-VL-4B-Instruct",
        "8b": "Qwen/Qwen3-VL-8B-Instruct",
    },
}

WHISPER_PROMPT = "Wykład akademicki w języku polskim. Poprawna interpunkcja i polskie znaki."

os.environ.setdefault("HF_HUB_DISABLE_SYMLINKS_WARNING", "1")
os.environ.setdefault("TOKENIZERS_PARALLELISM", "false")


# ============================================================================
# Wspólne
# ============================================================================

@dataclass
class Segment:
    start: float
    end: float
    text: str
    speaker: str | None = None


@dataclass
class Slide:
    index: int
    time: float
    file: str
    description: str | None = None
    material: str | None = None  # np. "wyklad.pdf, str. 5" – dopasowana strona prezentacji
    material_text: str | None = None


def fmt_ts(sec: float, millis: bool = True) -> str:
    sec = max(0.0, sec)
    h, m, s = int(sec // 3600), int(sec % 3600 // 60), sec % 60
    return f"{h:02d}:{m:02d}:{s:06.3f}" if millis else f"{h:02d}:{m:02d}:{int(s):02d}"


def fmt_dur(sec: float) -> str:
    sec = int(round(sec))
    if sec < 60:
        return f"{sec}s"
    if sec < 3600:
        return f"{sec // 60}m {sec % 60:02d}s"
    return f"{sec // 3600}h {sec % 3600 // 60:02d}m"


# ---------- etapy w stylu bootowania Linuksa ----------

_COLOR = sys.stdout.isatty() and os.environ.get("NO_COLOR") is None


def _c(code: str, s: str) -> str:
    return f"\033[{code}m{s}\033[0m" if _COLOR else s


class StageInfo:
    def __init__(self):
        self.note = ""
        self.skipped = False

    def skip(self, note: str) -> None:
        self.skipped, self.note = True, note


class Stages:
    """
             [2/9] Transkrypcja (Whisper)…
    [  OK  ] [2/9] Transkrypcja (Whisper) (4m 12s) — 812 segmentów
    """

    def __init__(self, total: int):
        self.total, self.n = total, 0

    @contextmanager
    def stage(self, name: str):
        self.n += 1
        tag = _c("1", f"[{self.n}/{self.total}]")
        print(f"         {tag} {name}…", flush=True)
        info, t0 = StageInfo(), time.time()
        try:
            yield info
        except Exception:
            print(f"[{_c('1;31', 'FAILED')}] {tag} {name}", flush=True)
            raise
        if info.skipped:
            print(f"[{_c('1;33', ' SKIP ')}] {tag} {name} — {info.note}", flush=True)
        else:
            note = f" — {info.note}" if info.note else ""
            print(f"[{_c('1;32', '  OK  ')}] {tag} {name} ({fmt_dur(time.time() - t0)}){note}", flush=True)


def err_msg(e: BaseException, limit: int = 160) -> str:
    """Pierwsza niepusta linia komunikatu błędu (albo nazwa typu, gdy komunikat jest pusty)."""
    lines = [ln.strip() for ln in str(e).splitlines() if ln.strip()]
    return (lines[0] if lines else type(e).__name__)[:limit]


def log(msg: str) -> None:
    print(f"           {msg}", flush=True)


# ---------- sprzęt / pamięć / pobieranie ----------

def is_apple_silicon() -> bool:
    return sys.platform == "darwin" and platform.machine() == "arm64"


def detect_device(requested: str = "auto") -> str:
    if requested != "auto":
        return requested
    try:
        import torch

        if torch.cuda.is_available():
            return "cuda"
    except Exception:
        pass
    try:
        import ctranslate2

        if ctranslate2.get_cuda_device_count() > 0:
            return "cuda"
    except Exception:
        pass
    return "mps" if is_apple_silicon() else "cpu"


def cuda_vram_gb() -> float:
    try:
        import torch

        return torch.cuda.get_device_properties(0).total_memory / 1024**3
    except Exception:
        return 0.0


def free_memory() -> None:
    gc.collect()
    try:
        import torch

        if torch.cuda.is_available():
            torch.cuda.empty_cache()
        if torch.backends.mps.is_available():
            torch.mps.empty_cache()
    except Exception:
        pass
    try:
        import mlx.core as mx

        mx.clear_cache()
    except Exception:
        pass


def hf_download(repo: str, target: Path, token: str | None = None) -> Path:
    """Pobiera repozytorium HF do ./models/... tylko jeśli go jeszcze nie ma."""
    from huggingface_hub import snapshot_download

    if not (target / ".done").exists():
        log(f"Pobieram {repo} → {target.relative_to(SCRIPT_DIR)} (jednorazowo)…")
        snapshot_download(repo, local_dir=str(target), token=token)
        (target / ".done").touch()
    return target


# ============================================================================
# 1. Transkrypcja (Whisper)
# ============================================================================

def load_audio(path: Path) -> np.ndarray:
    """Audio z dowolnego wideo/audio -> 16 kHz mono float32 (PyAV, bez ffmpeg w PATH).
    Własny dekoder zamiast faster_whisper.decode_audio – działa z każdą wersją PyAV
    (nowe wersje usunęły argument `metadata_errors`, którego używa faster-whisper)."""
    import av

    chunks = []
    with av.open(str(path)) as container:
        stream = next((s for s in container.streams if s.type == "audio"), None)
        if stream is None:
            raise RuntimeError("plik nie ma ścieżki dźwiękowej")
        stream.thread_type = "AUTO"
        resampler = av.AudioResampler(format="s16", layout="mono", rate=SAMPLE_RATE)
        for packet in container.demux(stream):
            try:
                frames = packet.decode()
            except Exception:  # uszkodzony pakiet – pomijamy, zamiast przerywać całe nagranie
                continue
            for frame in frames:
                for out in resampler.resample(frame):
                    chunks.append(out.to_ndarray().reshape(-1))
        for out in resampler.resample(None):  # opróżnienie bufora resamplera
            chunks.append(out.to_ndarray().reshape(-1))
    if not chunks:
        raise RuntimeError("nie udało się odczytać dźwięku z pliku")
    return np.concatenate(chunks).astype(np.float32) / 32768.0


def transcribe(audio: np.ndarray, model: str, language: str, device: str, prompt: str | None) -> list[Segment]:
    prompt = prompt or WHISPER_PROMPT
    if device == "mps":
        try:
            return _transcribe_mlx(audio, model, language, prompt)
        except ImportError:
            log("! mlx-whisper niedostępny – używam CPU")
            device = "cpu"
    if device == "cuda":
        try:
            return _transcribe_faster(audio, model, language, "cuda", prompt)
        except Exception as e:  # brak bibliotek CUDA 12, za mało VRAM, stary sterownik…
            msg = err_msg(e)
            log(f"! Whisper nie ruszył na GPU ({msg}) – przełączam na CPU")
            free_memory()
    return _transcribe_faster(audio, model, language, "cpu", prompt)


_CUDA_LIBS_LOADED = False


def _preload_cuda12_libs() -> None:
    """faster-whisper (CTranslate2) potrzebuje cuBLAS 12 i cuDNN 9. Ładujemy je z pakietów pip
    `nvidia-cublas-cu12` / `nvidia-cudnn-cu12`, żeby nie trzeba było ustawiać LD_LIBRARY_PATH –
    i żeby działało także wtedy, gdy PyTorch ma nowszą CUDA (np. 13) niż CTranslate2."""
    global _CUDA_LIBS_LOADED
    if _CUDA_LIBS_LOADED:
        return
    _CUDA_LIBS_LOADED = True
    import ctypes
    import glob

    roots = [os.path.join(p, "nvidia") for p in sys.path if os.path.isdir(os.path.join(p, "nvidia"))]
    if sys.platform == "win32":
        for root in roots:
            for sub in ("cublas", "cudnn"):
                d = os.path.join(root, sub, "bin")
                if os.path.isdir(d):
                    os.add_dll_directory(d)
                    os.environ["PATH"] = d + os.pathsep + os.environ.get("PATH", "")
        return
    patterns = ["cublas/lib/libcublasLt.so.12*", "cublas/lib/libcublas.so.12*", "cudnn/lib/libcudnn*.so.9*"]
    libs = [f for root in roots for pat in patterns for f in sorted(glob.glob(os.path.join(root, pat)))]
    # biblioteki cuDNN zależą od siebie nawzajem – ładujemy w kilku przebiegach
    pending = list(dict.fromkeys(libs))
    for _ in range(4):
        failed = []
        for lib in pending:
            try:
                ctypes.CDLL(lib, mode=ctypes.RTLD_GLOBAL)
            except OSError:
                failed.append(lib)
        if not failed or len(failed) == len(pending):
            break
        pending = failed


def _transcribe_faster(audio, model_name, language, device, prompt) -> list[Segment]:
    if device == "cuda":
        _preload_cuda12_libs()
    from faster_whisper import WhisperModel

    root = str(MODELS_DIR / "whisper")
    if device == "cuda":
        model = WhisperModel(model_name, device="cuda", compute_type="float16", download_root=root)
    else:
        model = WhisperModel(model_name, device="cpu", compute_type="int8",
                             cpu_threads=os.cpu_count() or 4, download_root=root)

    it, info = model.transcribe(
        audio,
        language=language,
        beam_size=5,
        vad_filter=True,
        vad_parameters={"min_silence_duration_ms": 500},
        initial_prompt=prompt,
        condition_on_previous_text=False,  # mniej zapętleń w długich nagraniach
    )
    out = []
    with tqdm(total=round(info.duration), unit="s", desc=f"           whisper ({device})", leave=False) as bar:
        for s in it:
            if s.text.strip():
                out.append(Segment(s.start, s.end, s.text.strip()))
            bar.update(max(0, round(s.end) - bar.n))
    del model
    free_memory()
    return out


def _transcribe_mlx(audio, model_name, language, prompt) -> list[Segment]:
    import mlx_whisper

    repo = MLX_WHISPER.get(model_name, model_name)
    local = hf_download(repo, MODELS_DIR / "whisper-mlx" / repo.split("/")[-1])
    result = mlx_whisper.transcribe(
        audio,
        path_or_hf_repo=str(local),
        language=language,
        initial_prompt=prompt,
        condition_on_previous_text=False,
        verbose=False,  # pasek postępu
    )
    free_memory()
    return [Segment(float(s["start"]), float(s["end"]), s["text"].strip())
            for s in result["segments"] if s["text"].strip()]


# ============================================================================
# 2. Mówcy (pyannote) — tylko z --speakers
# ============================================================================

def _sha256(path: Path) -> str:
    import hashlib

    h = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def _get_pyannote(token: str | None) -> Path:
    """Pobiera model mówców do ./models (bez tokenu) i sprawdza integralność plików."""
    import shutil

    from huggingface_hub import snapshot_download

    target = MODELS_DIR / "pyannote-community-1"
    if (target / ".done").exists():
        return target

    def verified() -> bool:
        return all((target / f).exists() and _sha256(target / f) == h for f, h in PYANNOTE_SHA256.items())

    sources = [(PYANNOTE_MIRROR, PYANNOTE_REVISION, None)]
    if token:  # zapasowo: oficjalne repo, jeśli ktoś ma token
        sources.append((PYANNOTE_REPO, None, token))
    last_err: Exception | None = None
    for repo, rev, tok in sources:
        try:
            log(f"Pobieram {repo} → {target.relative_to(SCRIPT_DIR)} (jednorazowo, ~33 MB)…")
            snapshot_download(repo, revision=rev, local_dir=str(target), token=tok)
            if verified():
                (target / ".done").touch()
                return target
            last_err = RuntimeError(f"pliki z {repo} mają inne sumy kontrolne niż oczekiwane")
        except Exception as e:
            last_err = e
        shutil.rmtree(target, ignore_errors=True)
    raise RuntimeError(f"nie udało się pobrać modelu mówców: {last_err}")


def diarize(audio: np.ndarray, device: str, token: str | None,
            num: int | None, lo: int | None, hi: int | None) -> list[tuple[float, float, str]]:
    try:
        import torch
        from pyannote.audio import Pipeline
    except ImportError as e:
        raise RuntimeError("brak pakietu pyannote.audio (pip install pyannote.audio)") from e

    local = _get_pyannote(token or os.environ.get("HF_TOKEN"))
    pipeline = Pipeline.from_pretrained(str(local))

    torch_dev = device if device in ("cuda", "mps") else "cpu"
    if torch_dev == "cuda" and not torch_cuda_ok()[0]:
        torch_dev = "cpu"
    try:
        pipeline.to(torch.device(torch_dev))
    except Exception:
        torch_dev = "cpu"

    data = {"waveform": torch.from_numpy(np.ascontiguousarray(audio)).float().unsqueeze(0),
            "sample_rate": SAMPLE_RATE}
    kwargs = {k: v for k, v in {"num_speakers": num, "min_speakers": lo, "max_speakers": hi}.items() if v}

    def run():
        from pyannote.audio.pipelines.utils.hook import ProgressHook

        with ProgressHook(transient=True) as hook:
            return pipeline(data, hook=hook, **kwargs)

    try:
        output = run()
    except RuntimeError as e:
        if torch_dev != "mps":
            raise
        log(f"! MPS: {e} – ponawiam na CPU")
        pipeline.to(torch.device("cpu"))
        output = run()

    annotation = output  # pyannote 4: DiarizeOutput; exclusive = bez nakładek (lepsze do transkrypcji)
    for attr in ("exclusive_speaker_diarization", "speaker_diarization"):
        if getattr(output, attr, None) is not None:
            annotation = getattr(output, attr)
            break
    turns = [(t.start, t.end, spk) for t, _, spk in annotation.itertracks(yield_label=True)]
    del pipeline
    free_memory()
    return turns


def assign_speakers(segments: list[Segment], turns: list[tuple[float, float, str]]) -> list[Segment]:
    """Mówca o największym pokryciu czasowym segmentu; etykiety -> 'Mówca N' wg kolejności."""
    if not turns:
        return segments
    turns = sorted(turns)
    for seg in segments:
        best, best_ov = None, 0.0
        for s, e, spk in turns:
            if s > seg.end:
                break
            ov = min(seg.end, e) - max(seg.start, s)
            if ov > best_ov:
                best, best_ov = spk, ov
        if best is None:
            mid = (seg.start + seg.end) / 2
            best = min(turns, key=lambda t: min(abs(t[0] - mid), abs(t[1] - mid)))[2]
        seg.speaker = best
    names: dict[str, str] = {}
    for seg in segments:
        names.setdefault(seg.speaker, f"Mówca {len(names) + 1}")
        seg.speaker = names[seg.speaker]
    return segments


# ============================================================================
# 3. Slajdy (OpenCV)
# ============================================================================

def _prep(frame):
    """Klatka -> (szarość 640px, mapa krawędzi = 'treść' slajdu: tekst, linie, wykresy)."""
    import cv2

    h, w = frame.shape[:2]
    gray = cv2.cvtColor(cv2.resize(frame, (640, max(1, int(h * 640 / w))), interpolation=cv2.INTER_AREA),
                        cv2.COLOR_BGR2GRAY)
    edges = cv2.Canny(cv2.GaussianBlur(gray, (3, 3), 0), 60, 160)
    # pogrubienie krawędzi = odporność na drobne przesunięcia/kompresję
    return gray, cv2.dilate(edges, np.ones((3, 3), np.uint8)) > 0


def _change(a, b, static) -> float:
    """Jaka część TREŚCI kadru się zmieniła (0..1), tylko w statycznych obszarach.
    Liczone względem treści, nie całego ekranu – dzięki temu zmiana slajdu z samym tekstem
    na białym tle daje wysoką wartość, a drgania/kompresja – bliską zera."""
    a, b = a & static, b & static
    union = int((a | b).sum())
    return float((a ^ b).sum() / union) if union > 200 else 0.0


def extract_slides(video: Path, out_dir: Path, interval: float = 1.0, threshold: float = 0.35,
                   min_stable: int = 2, max_width: int = 1920) -> list[Slide]:
    """
    Zapisuje klatkę tylko przy ISTOTNEJ zmianie treści:
      • obraz musi być stabilny przez `min_stable` próbek (pomija przejścia i animacje),
      • co najmniej `threshold` treści (tekst, linie, wykresy) musi być nowe względem
        poprzedniego slajdu – drobne dopiski i podkreślenia nie tworzą nowego slajdu,
      • obszary z ciągłym ruchem (kamerka, prowadzący) są wykrywane i ignorowane,
      • powrót do slajdu, który już był, nie tworzy duplikatu.
    """
    import cv2

    out_dir.mkdir(parents=True, exist_ok=True)
    for f in out_dir.glob("slajd_*.jpg"):
        f.unlink()
    cap = cv2.VideoCapture(str(video))
    if not cap.isOpened():
        raise RuntimeError(f"Nie można otworzyć wideo: {video}")
    fps = cap.get(cv2.CAP_PROP_FPS) or 25.0
    total = int(cap.get(cv2.CAP_PROP_FRAME_COUNT)) or None
    step = max(1, round(fps * interval))
    stable_eps = min(0.1, threshold / 3)  # poniżej tego obraz uznajemy za "stojący"

    slides: list[Slide] = []
    saved: list = []  # mapy treści zapisanych slajdów (do wykrywania powtórek)
    prev = last = stable_frame = activity = None
    stable, stable_since, idx, dupes, updates = 0, 0.0, 0, 0, 0

    with tqdm(total=total, unit="kl", unit_scale=True, desc="           klatki", leave=False) as bar:
        while cap.grab():
            idx += 1
            bar.update(1)
            if (idx - 1) % step:
                continue
            ok, frame = cap.retrieve()
            if not ok:
                continue
            t = (idx - 1) / fps
            gray, edges = _prep(frame)

            if prev is None:
                activity = np.zeros(edges.shape, np.float32)
                static = np.ones(edges.shape, bool)
                moving = 1.0
            else:
                # mapa aktywności: jednorazowa zmiana slajdu ledwo ją podnosi,
                # ciągły ruch (kamerka, prowadzący) szybko dochodzi do ~1 i zostaje zamaskowany
                activity = 0.85 * activity + 0.15 * (edges ^ prev)
                static = cv2.erode((activity < 0.2).astype(np.uint8), np.ones((9, 9), np.uint8)) > 0
                moving = _change(edges, prev, static)
            prev = edges

            if moving < stable_eps:
                stable += 1
            else:
                stable, stable_since, stable_frame = 0, t, frame

            if stable + 1 < min_stable or gray.std() < 4:  # niestabilny albo czarny/jednolity
                continue
            if last is not None and _change(edges, last, static) <= threshold:
                # ten sam slajd; jeśli coś na nim DOPISANO (kolejny punkt, wzór), podmieniamy
                # zapisany obraz na pełniejszą wersję – opis obejmie wtedy całą treść
                if slides and last is saved[-1] and _change(edges, last, static) > stable_eps \
                        and (edges & static).sum() > (last & static).sum() * 1.03:
                    img = stable_frame
                    if img.shape[1] > max_width:
                        img = cv2.resize(img, (max_width, int(img.shape[0] * max_width / img.shape[1])),
                                         interpolation=cv2.INTER_AREA)
                    cv2.imwrite(str(out_dir / slides[-1].file), img, [cv2.IMWRITE_JPEG_QUALITY, 92])
                    saved[-1] = last = edges
                    updates += 1
                continue
            last = edges
            if any(_change(edges, e, static) <= threshold for e in saved):
                dupes += 1  # powrót do slajdu, który już był
                continue

            i = len(slides) + 1
            img = stable_frame
            if img.shape[1] > max_width:
                img = cv2.resize(img, (max_width, int(img.shape[0] * max_width / img.shape[1])),
                                 interpolation=cv2.INTER_AREA)
            name = f"slajd_{i:03d}_{fmt_ts(stable_since, False).replace(':', '-')}.jpg"
            cv2.imwrite(str(out_dir / name), img, [cv2.IMWRITE_JPEG_QUALITY, 92])
            slides.append(Slide(i, stable_since, name))
            saved.append(edges)
    cap.release()
    if dupes:
        log(f"pominięto powroty do wcześniejszych slajdów: {dupes}")
    if updates:
        log(f"uzupełniono slajdy o dopisaną treść: {updates}")
    return slides


# ============================================================================
# 4. Qwen3-VL — czysty Python (mlx-vlm na Macu, transformers na NVIDIA/CPU)
# ============================================================================

_THINK = re.compile(r"<think>.*?</think>", re.DOTALL)


def _vlm_image(path: Path, max_side: int = 1280) -> str:
    """Zmniejsza slajd do rozsądnej wielkości (mniej tokenów = szybciej), zwraca ścieżkę tymczasową."""
    import cv2

    img = cv2.imread(str(path))
    h, w = img.shape[:2]
    if max(h, w) > max_side:
        k = max_side / max(h, w)
        img = cv2.resize(img, (int(w * k), int(h * k)), interpolation=cv2.INTER_AREA)
    fd, tmp = tempfile.mkstemp(suffix=".jpg")
    os.close(fd)
    cv2.imwrite(tmp, img, [cv2.IMWRITE_JPEG_QUALITY, 90])
    return tmp


def torch_cuda_ok() -> tuple[bool, str]:
    """Czy PyTorch naprawdę może liczyć na karcie (nie tylko czy ją 'widzi')."""
    try:
        import torch
    except ImportError:
        return False, "brak PyTorcha"
    if not torch.cuda.is_available():
        return False, f"torch {torch.__version__} (CUDA {torch.version.cuda or 'brak – wersja CPU'}) nie widzi karty"
    try:
        (torch.ones(8, device="cuda") * 2).sum().item()
        return True, ""
    except Exception as e:
        return False, f"torch nie może liczyć na karcie: {err_msg(e)}"


class Qwen:
    def __init__(self, size: str, device: str):
        self.backend = {"mps": "mlx", "cuda": "cuda"}.get(device, "cpu")
        if self.backend == "mlx":
            try:
                import mlx_vlm  # noqa: F401
            except ImportError:
                log("! mlx-vlm niedostępny – używam CPU")
                self.backend = "cpu"
        if self.backend == "cuda":
            ok, why = torch_cuda_ok()
            if not ok:
                log(f"! Qwen nie może użyć karty: {why} – używam CPU")
                self.backend = "cpu"

        custom = size if "/" in size else None
        if not custom and size == "auto":
            size = {"mlx": "8b", "cpu": "4b"}.get(self.backend) or ("8b" if cuda_vram_gb() >= 7.5 else "4b")
        self.size = size

        if self.backend == "mlx":
            self._load_mlx(custom or QWEN["mlx"][size])
            return

        # kolejne próby: każda sprawdzana krótkim testem generowania
        attempts = []
        if self.backend == "cuda":
            if custom:
                attempts.append(("cuda", custom, False))
            else:
                attempts.append(("cuda", QWEN["cuda"][size], False))  # prekwantyzowany 4-bit (mały download)
                attempts.append(("cuda", QWEN["cpu"][size], True))    # oryginał + kwantyzacja 4-bit w locie
        cpu_size = "4b" if size == "8b" else size  # 8B w bf16 nie zmieści się w 16 GB RAM
        attempts.append(("cpu", custom or QWEN["cpu"][cpu_size], False))

        import transformers

        def marker(repo: str) -> Path:  # "ten model nie działa z tą wersją transformers"
            return MODELS_DIR / "qwen" / repo.split("/")[-1] / f".nie_dziala_tf{transformers.__version__}"

        if len(attempts) > 1 and marker(attempts[0][1]).exists():
            attempts = attempts[1:]  # wiemy z poprzedniego uruchomienia, że to nie działa
        last: Exception | None = None
        for i, (backend, repo, quantize) in enumerate(attempts):
            try:
                self._load_hf(backend, repo, quantize)
                self._smoke_test()
                return
            except Exception as e:
                last = e
                self._unload()
                if i + 1 < len(attempts) and not quantize and backend == "cuda":
                    try:
                        marker(repo).touch()
                    except OSError:
                        pass
                if i + 1 < len(attempts):
                    nb, nr, nq = attempts[i + 1]
                    how = "4-bit w locie" if nq else ("CPU" if nb == "cpu" else nb)
                    log(f"! {repo.split('/')[-1]} nie działa ({err_msg(e)})")
                    log(f"  próbuję: {nr.split('/')[-1]} ({how})")
        raise RuntimeError(f"nie udało się uruchomić Qwen3-VL: {last}")

    def _load_mlx(self, repo: str) -> None:
        from mlx_vlm import load
        from mlx_vlm.utils import load_config

        self.name = repo.split("/")[-1]
        path = hf_download(repo, MODELS_DIR / "qwen" / self.name)
        log(f"Ładuję {self.name} (mlx)…")
        self.model, self.processor = load(str(path))
        self.config = load_config(str(path))

    def _load_hf(self, backend: str, repo: str, quantize: bool) -> None:
        import torch
        from transformers import AutoModelForImageTextToText, AutoProcessor

        self.backend = backend
        self.name = repo.split("/")[-1] + (" 4-bit" if quantize else "")
        path = hf_download(repo, MODELS_DIR / "qwen" / repo.split("/")[-1])
        log(f"Ładuję {self.name} ({backend})…")
        kw = {"dtype": torch.bfloat16, "device_map": "cuda" if backend == "cuda" else "cpu"}
        if quantize:
            from transformers import BitsAndBytesConfig

            kw["quantization_config"] = BitsAndBytesConfig(
                load_in_4bit=True, bnb_4bit_quant_type="nf4",
                bnb_4bit_compute_dtype=torch.bfloat16, bnb_4bit_use_double_quant=True,
                llm_int8_skip_modules=["visual", "lm_head"],  # wieża wizyjna w pełnej precyzji
            )
        if backend == "cpu":
            torch.set_num_threads(os.cpu_count() or 4)
        self.model = AutoModelForImageTextToText.from_pretrained(str(path), **kw).eval()
        self.processor = AutoProcessor.from_pretrained(str(path))

    def _smoke_test(self) -> None:
        """Krótki test Z OBRAZKIEM – sprawdza też część wizyjną modelu (to ona psuje się przy
        niezgodnych wersjach kwantyzacji), więc błąd wychodzi od razu, a nie przy pierwszym slajdzie."""
        import cv2

        img = np.full((224, 448, 3), 255, np.uint8)
        cv2.putText(img, "TEST", (100, 145), cv2.FONT_HERSHEY_SIMPLEX, 3, (0, 0, 0), 8)
        fd, tmp = tempfile.mkstemp(suffix=".png")
        os.close(fd)
        try:
            cv2.imwrite(tmp, img)
            out = self.chat("Jaki napis jest na obrazku? Odpowiedz jednym słowem.", image=Path(tmp), max_tokens=4,
                            max_rounds=0)
        finally:
            os.unlink(tmp)
        if not out:
            raise RuntimeError("model nie zwrócił odpowiedzi")

    def _unload(self) -> None:
        for attr in ("model", "processor", "config"):
            if hasattr(self, attr):
                delattr(self, attr)
        free_memory()

    def _generate(self, prompt: str, image: Path | None, max_tokens: int, progress: str | None,
                  prefix: str = "") -> str:
        """Surowa odpowiedź modelu; `prefix` = początek odpowiedzi, którą model ma kontynuować.
        Ustawia self.truncated, gdy generowanie zatrzymał limit tokenów (a nie koniec odpowiedzi)."""
        img = _vlm_image(image) if image else None
        bar = tqdm(desc=progress, unit=" tok", leave=False) if progress else None
        self.truncated = False
        try:
            parts = []
            gen = self._mlx if self.backend == "mlx" else self._hf
            for piece in gen(prompt, img, max_tokens, prefix):
                parts.append(piece)
                if bar is not None:
                    bar.update(1)
        finally:
            if bar is not None:
                bar.close()
            if img:
                os.unlink(img)
        return "".join(parts)

    def chat(self, prompt: str, image: Path | None = None, max_tokens: int = 2048,
             progress: str | None = None, max_rounds: int = 3) -> str:
        """Odpowiedź modelu. Gdy urwie się na limicie tokenów, model dostaje swój dotychczasowy tekst
        jako początek odpowiedzi i pisze dalej (do `max_rounds` razy) – nic nie zostaje ucięte w połowie."""
        text = self._generate(prompt, image, max_tokens, progress)
        rounds = 0
        while self.truncated and rounds < max_rounds:
            rounds += 1
            text += self._generate(prompt, image, max_tokens,
                                   f"{progress} (cd. {rounds})" if progress else None, prefix=text)
        if self.truncated:
            # wciąż za długo – obcinamy do ostatniego pełnego akapitu, zamiast zostawiać pół zdania
            cut = text.rstrip().rfind("\n\n")
            if cut < len(text) * 0.6:  # brak akapitu blisko końca – tniemy do ostatniego pełnego zdania
                cut = max(text.rfind(". "), text.rfind(".\n"), text.rfind("? "), text.rfind("! ")) + 1
            if cut > len(text) * 0.6:
                text = text[:cut]
        return _THINK.sub("", text).strip()

    def _mlx(self, prompt, img, max_tokens, prefix=""):
        from mlx_vlm import stream_generate
        from mlx_vlm.prompt_utils import apply_chat_template

        formatted = apply_chat_template(self.processor, self.config, prompt, num_images=1 if img else 0)
        formatted += prefix  # kontynuacja: dopisujemy początek odpowiedzi asystenta
        last = None
        for r in stream_generate(self.model, self.processor, formatted, image=[img] if img else None,
                                 max_tokens=max_tokens, temperature=0.0, repetition_penalty=1.05):
            last = r
            yield r.text
        if last is not None:
            self.truncated = getattr(last, "finish_reason", None) == "length" or \
                getattr(last, "generation_tokens", 0) >= max_tokens

    def _hf(self, prompt, img, max_tokens, prefix=""):
        import torch
        from transformers import TextIteratorStreamer

        content = ([{"type": "image", "image": img}] if img else []) + [{"type": "text", "text": prompt}]
        messages = [{"role": "user", "content": content}]
        if prefix:
            messages.append({"role": "assistant", "content": [{"type": "text", "text": prefix}]})
        inputs = self.processor.apply_chat_template(
            messages, tokenize=True, return_dict=True, return_tensors="pt",
            add_generation_prompt=not prefix, continue_final_message=bool(prefix),
        ).to(self.model.device)
        n_in = inputs["input_ids"].shape[1]
        streamer = TextIteratorStreamer(self.processor.tokenizer, skip_prompt=True, skip_special_tokens=True)
        errors, n_out = [], [0]

        def gen():
            try:
                with torch.inference_mode():
                    out = self.model.generate(**inputs, max_new_tokens=max_tokens, do_sample=False,
                                              repetition_penalty=1.05, streamer=streamer)
                n_out[0] = out.shape[1] - n_in
            except Exception as e:  # przekaż błąd do wątku głównego
                errors.append(e)
                streamer.end()

        th = threading.Thread(target=gen, daemon=True)
        th.start()
        yield from streamer
        th.join()
        if errors:
            raise errors[0]
        self.truncated = n_out[0] >= max_tokens

    def close(self) -> None:
        self._unload()


SLIDE_PROMPT = """Opisz po polsku ten slajd z wykładu, tak aby student mógł się z niego uczyć bez patrzenia na obraz.
- Przepisz dokładnie tytuł i wszystkie punkty tekstu.
- Wzory zapisz w LaTeX ($...$).
- Opisz diagramy, wykresy, tabele i kod (kod przepisz w bloku ```).
- Jeśli to nie slajd (np. sama kamera z prowadzącym), napisz to w jednym zdaniu.
Nie dodawaj informacji, których nie ma na obrazie. Odpowiedz samym opisem."""

CHUNK_PROMPT = """Poniżej fragment transkrypcji wykładu ({start}–{end}) z opisami slajdów wyświetlanych w tym czasie.
Transkrypcja pochodzi z automatycznego rozpoznawania mowy – może zawierać błędy, popraw je z kontekstu.

Sporządź szczegółowe notatki do nauki z tego fragmentu w Markdown:
- definicje, twierdzenia, wzory (LaTeX), przykłady, polecenia/kod i ważne uwagi prowadzącego.

Format – ściśle:
- każde zagadnienie jako nagłówek `### Tytuł [hh:mm:ss]` (czas rozpoczęcia); podpunkty ewentualnie jako `####`,
- NIE używaj nagłówków `#` ani `##`,
- NIE dodawaj sekcji „Podsumowanie”, „Pytania”, „Pojęcia”, „TL;DR” ani „Slajdy” – zostaną dodane osobno,
- o slajdzie pisz tam, gdzie prowadzący go omawia, z dopiskiem (slajd N) w tekście –
  nie wypisuj slajdów w osobnej liście.
Jeśli przy slajdzie jest „Treść z materiałów”, to dokładny tekst slajdu z prezentacji prowadzącego –
przy definicjach i wzorach trzymaj się go (opis z nagrania może zawierać błędy).
Pisz po polsku, rzeczowo, bez wstępów.

--- FRAGMENT ---
{body}"""

HEAD_PROMPT = """Poniżej notatki z kolejnych części jednego wykładu.
Napisz WYŁĄCZNIE początek notatki do nauki w Markdown, dokładnie w tym formacie:

# <krótki, konkretny tytuł wykładu>

## TL;DR
<3–5 zdań: o czym był wykład i co z niego najważniejsze>

## Najważniejsze pojęcia
- **<pojęcie>** – <krótka definicja>
<8–15 pozycji, w kolejności z wykładu>

Nie pisz nic więcej (żadnych innych sekcji). Pisz po polsku, rzeczowo, nie wymyślaj faktów.
{web}
--- NOTATKI ---
{body}"""

QUESTIONS_PROMPT = """Na podstawie poniższych notatek z wykładu ułóż 6–10 pytań kontrolnych sprawdzających zrozumienie
materiału – od prostszych do trudniejszych. Wypisz je jako listę numerowaną, bez odpowiedzi, bez nagłówka
i bez wstępu. Po polsku.

--- NOTATKI ---
{body}"""

CONDENSE_PROMPT = """Streść poniższe notatki z części wykładu w zwięzłych punktach (maks. ok. 250 słów).
Zachowaj wszystkie kluczowe pojęcia, definicje i wzory. Bez wstępów, po polsku.

{body}"""


def describe_slides(qwen: Qwen, slides: list[Slide], slides_dir: Path) -> None:
    for s in tqdm([s for s in slides if not s.description], desc="           qwen-vl", unit="slajd", leave=False):
        s.description = qwen.chat(SLIDE_PROMPT, image=slides_dir / s.file, max_tokens=1536, max_rounds=1)


def timeline_chunks(segments: list[Segment], slides: list[Slide], max_chars: int):
    """Przeplata transkrypcję z opisami slajdów (po czasie) i tnie na fragmenty."""
    events = [(s.start, 1, f"[{fmt_ts(s.start, False)}] {s.speaker + ': ' if s.speaker else ''}{s.text}")
              for s in segments]
    for sl in slides:
        if not sl.description and not sl.material_text:
            continue
        txt = f"\n[SLAJD {sl.index} @ {fmt_ts(sl.time, False)}]\n"
        if sl.material_text:
            txt += f"Treść z materiałów ({sl.material}):\n{sl.material_text}\n"
            txt += f"Opis z nagrania:\n{sl.description}\n" if sl.description else ""
        else:
            txt += f"{sl.description}\n"
        events.append((sl.time, 0, txt))
    events.sort(key=lambda e: (e[0], e[1]))

    chunks, buf, size, start = [], [], 0, None
    for t, _, line in events:
        start = t if start is None else start
        if size + len(line) > max_chars and buf:
            chunks.append((start, t, "\n".join(buf)))
            buf, size, start = [], 0, t
        buf.append(line)
        size += len(line) + 1
    if buf:
        chunks.append((start or 0, segments[-1].end if segments else 0, "\n".join(buf)))
    return chunks


MAX_CHARS = 12000  # ~4-5k tokenów na fragment – bezpiecznie dla 16 GB RAM


def partial_notes(qwen: Qwen, segments: list[Segment], slides: list[Slide],
                  pages: list[Page] | None = None, cache: Path | None = None) -> tuple[list[str], int]:
    """Notatki z kolejnych fragmentów. Każda gotowa notatka trafia od razu do `cache`, więc po
    przerwaniu/błędzie kolejne uruchomienie liczy tylko brakujące fragmenty (te same dane = ten sam klucz)."""
    import hashlib

    jobs = [CHUNK_PROMPT.format(start=fmt_ts(a, False), end=fmt_ts(b, False), body=body)
            for a, b, body in timeline_chunks(segments, slides, MAX_CHARS)]
    n_lecture = len(jobs)
    jobs += [MATERIAL_PROMPT.format(body=body) for body in material_chunks(pages or [], MAX_CHARS)]

    def key(prompt: str) -> str:
        return hashlib.sha256(f"v2\n{qwen.name}\n{prompt}".encode()).hexdigest()[:24]

    saved: dict[str, str] = {}
    if cache and cache.exists():
        try:
            saved = json.loads(cache.read_text("utf-8"))
        except Exception:
            saved = {}
    out, reused = [], 0
    for i, prompt in enumerate(tqdm(jobs, desc="           fragmenty", unit="fragm", leave=False)):
        k = key(prompt)
        if k in saved:
            note, reused = saved[k], reused + 1
        else:
            note = qwen.chat(prompt, max_tokens=3072)
            saved[k] = note
            if cache:
                cache.write_text(json.dumps(saved, ensure_ascii=False), "utf-8")
        out.append(note if i < n_lecture else f"{MATERIAL_TAG}\n{note}")
    return out, reused


MATERIAL_TAG = "[Z MATERIAŁÓW]"


_DROP_SECTIONS = re.compile(
    r"^(tl;?dr|pytania( kontrolne| sprawdzające| do powtórki)?|(najważniejsze|kluczowe) pojęcia|"
    r"słowniczek|podsumowanie( fragmentu)?)\s*[:.]?\s*$", re.IGNORECASE)
_DROP_HEADING_ONLY = re.compile(r"^(przebieg wykładu|notatk\w*( z \w+)?|fragment\w*)\b", re.IGNORECASE)


def _as_subsections(note: str) -> str:
    """Normalizuje notatkę z fragmentu, żeby po sklejeniu struktura dokumentu była spójna:
      • usuwa sekcje dublujące części globalne (TL;DR, pytania, pojęcia, podsumowanie fragmentu),
      • usuwa nagłówki-meta ("Notatki z fragmentu…", "Przebieg wykładu"), zostawiając treść,
      • przesuwa poziomy nagłówków tak, że najwyższy to ### (pod "## Przebieg wykładu"), maks. #####."""
    lines = note.strip().splitlines()
    heads, fence = [], False
    for i, line in enumerate(lines):
        st = line.strip()
        if st.startswith("```"):
            fence = not fence
        m = None if fence else _HEADING.match(line)
        if m:
            title = re.sub(r"\s*\[[\d:–\-\s]+\]\s*$", "", m.group(2).strip()).strip("*_ ")
            heads.append((i, len(m.group(1)), m.group(2).strip(), title))

    drop: set[int] = set()
    for k, (i, lvl, _, title) in enumerate(heads):
        if _DROP_SECTIONS.match(title):
            end = next((j for j, l2, _, _ in heads[k + 1:] if l2 <= lvl), len(lines))
            drop.update(range(i, end))
        elif _DROP_HEADING_ONLY.match(title):
            drop.add(i)
    kept = [(i, lvl, txt) for i, lvl, txt, _ in heads if i not in drop]
    base = min((lvl for _, lvl, _ in kept), default=3)
    new_head = {i: "#" * min(5, 3 + lvl - base) + " " + txt for i, lvl, txt in kept}
    out = [new_head.get(i, line) for i, line in enumerate(lines) if i not in drop]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(out)).strip()


def _strip_leading_headings(text: str) -> str:
    lines = text.strip().splitlines()
    while lines and (_HEADING.match(lines[0]) or not lines[0].strip()):
        lines.pop(0)
    return "\n".join(lines).strip()


def condense(qwen: Qwen, notes: list[str], limit: int = 2 * MAX_CHARS) -> str:
    """Skrót wszystkich notatek, mieszczący się w kontekście – do tytułu, TL;DR i pytań."""
    while sum(map(len, notes)) > limit and notes:
        groups, cur, size = [], [], 0
        for n in notes:
            if size + len(n) > MAX_CHARS and cur:
                groups.append(cur)
                cur, size = [], 0
            cur.append(n)
            size += len(n)
        groups.append(cur)
        notes = [qwen.chat(CONDENSE_PROMPT.format(body="\n\n".join(g)), max_tokens=768)
                 for g in tqdm(groups, desc="           streszczanie", unit="część", leave=False)]
        if len(groups) == 1:
            break
    return "\n\n---\n\n".join(notes)[: limit + MAX_CHARS]


def final_notes(qwen: Qwen, partial: list[str], web: str = "") -> str:
    """Składa notatkę końcową z części, zamiast generować całość jednym (długim, ucinanym) wywołaniem:
      • tytuł + TL;DR + pojęcia  – krótkie wywołanie modelu na skrócie całości,
      • Przebieg wykładu          – szczegółowe notatki z kolejnych fragmentów (już policzone),
      • Z materiałów              – notatki ze stron prezentacji nieomówionych na nagraniu,
      • Pytania kontrolne / Dodatkowe materiały – osobne krótkie wywołania."""
    lecture = [p for p in partial if not p.startswith(MATERIAL_TAG)]
    mats = [p[len(MATERIAL_TAG):].strip() for p in partial if p.startswith(MATERIAL_TAG)]
    digest = condense(qwen, lecture + mats)

    head = qwen.chat(HEAD_PROMPT.format(body=digest, web=WEB_HEAD.format(web=web) if web else ""),
                     max_tokens=2048, progress="           tytuł, TL;DR, pojęcia")
    if not head.lstrip().startswith("# "):
        head = "# Notatka z wykładu\n\n" + head
    first, rest = head.strip().split("\n", 1) if "\n" in head.strip() else (head.strip(), "")
    rest = re.sub(r"^#\s", "## ", rest, flags=re.MULTILINE)  # kolejne "# …" w głowie -> "## …"
    head = first + ("\n" + rest if rest else "")
    questions = _strip_leading_headings(
        qwen.chat(QUESTIONS_PROMPT.format(body=digest), max_tokens=1024, progress="           pytania kontrolne"))
    sources = ""
    if web:
        topic = re.sub(r"\s+", " ", head)[:600]
        sources = _strip_leading_headings(
            qwen.chat(WEB_SECTION_PROMPT.format(web=web, topic=topic), max_tokens=1536,
                      progress="           dodatkowe materiały"))

    parts = [head.strip(), "## Przebieg wykładu", *(_as_subsections(n) for n in lecture)]
    if mats:
        parts += ["## Z materiałów (nieomówione na wykładzie)", *(_as_subsections(n) for n in mats)]
    if questions:
        parts += ["## Pytania kontrolne", questions]
    if sources:
        parts += ["## Dodatkowe materiały", sources]
    return "\n\n".join(p for p in parts if p.strip()) + "\n"


# ============================================================================
# 4b. Materiały dodatkowe (--materials): prezentacje, notatki, obrazy
# ============================================================================

MATERIAL_EXTS = {".pdf", ".pptx", ".docx", ".txt", ".md", ".png", ".jpg", ".jpeg", ".webp"}
AUTO_MATERIAL_EXTS = (".pdf", ".pptx")  # wyklad.pdf obok wyklad.mp4 dołączamy automatycznie

MATERIAL_PROMPT = """Poniżej fragmenty materiałów do wykładu (prezentacja / notatki prowadzącego), których nie udało się
dopasować do konkretnego momentu nagrania. Sporządź z nich zwięzłe notatki do nauki w Markdown:
definicje, twierdzenia, wzory (LaTeX), przykłady. Przy każdej sekcji podaj źródło w nawiasie, np. (wyklad.pdf, str. 5).
Pisz po polsku, rzeczowo, bez wstępów.

--- MATERIAŁY ---
{body}"""


@dataclass
class Page:
    doc: str
    num: int
    text: str
    image: str | None = None  # wyrenderowana strona / obraz – do opisania modelem
    description: str | None = None
    matched: bool = False

    @property
    def ref(self) -> str:
        return f"{self.doc}, str. {self.num}"


def _clean(text: str) -> str:
    lines = (re.sub(r"[ \t\u00a0]+", " ", ln).strip() for ln in text.replace("\r", "\n").splitlines())
    return "\n".join(ln for ln in lines if ln)


def _split(text: str, size: int = 3000) -> list[str]:
    parts, cur = [], ""
    for para in text.split("\n"):
        if len(cur) + len(para) > size and cur:
            parts.append(cur)
            cur = ""
        cur += para + "\n"
    return parts + ([cur] if cur.strip() else [])


def _pptx_text(shapes) -> list[str]:
    out = []
    for sh in shapes:
        if getattr(sh, "shape_type", None) == 6:  # grupa kształtów
            out += _pptx_text(sh.shapes)
        elif getattr(sh, "has_text_frame", False) and sh.has_text_frame:
            out.append(sh.text_frame.text)
        elif getattr(sh, "has_table", False) and sh.has_table:
            out.append("\n".join(" | ".join(c.text for c in row.cells) for row in sh.table.rows))
    return out


def _pdf_image_share(page) -> float:
    """Jaką część strony zajmują obrazy (wykresy, zdjęcia, zrzuty) – wtedy warto ją opisać modelem VL."""
    try:
        import pypdfium2.raw as raw

        w, h = page.get_size()
        area = 0.0
        for obj in page.get_objects(filter=(raw.FPDF_PAGEOBJ_IMAGE,)):
            left, bottom, right, top = obj.get_bounds()
            area += max(0.0, right - left) * max(0.0, top - bottom)
        return min(1.0, area / (w * h)) if w and h else 0.0
    except Exception:
        return 0.0


def load_materials(files: list[Path], render_dir: Path) -> list[Page]:
    """PDF/PPTX -> strona po stronie; DOCX/TXT/MD -> kawałki tekstu; obrazy -> do opisania.
    Strony PDF bez tekstu (skany, diagramy) albo z dużym obrazem są renderowane i opisywane przez Qwen3-VL;
    resztę bierzemy jako czysty tekst – szybko i bez błędów odczytu."""
    pages: list[Page] = []
    for f in files:
        ext = f.suffix.lower()
        if ext == ".pdf":
            import pypdfium2 as pdfium

            pdf = pdfium.PdfDocument(str(f))
            for i in range(len(pdf)):
                page = pdf[i]
                text = _clean(page.get_textpage().get_text_range())
                img = None
                if len(text) < 40 or _pdf_image_share(page) > 0.25:  # skan / diagram / duży obraz
                    render_dir.mkdir(parents=True, exist_ok=True)
                    img = str(render_dir / f"{f.stem}_str{i + 1:03d}.png")
                    page.render(scale=2).to_pil().save(img)
                pages.append(Page(f.name, i + 1, text, img))
            pdf.close()
        elif ext == ".pptx":
            from pptx import Presentation

            for i, slide in enumerate(Presentation(str(f)).slides, 1):
                parts = _pptx_text(slide.shapes)
                if slide.has_notes_slide:
                    notes = slide.notes_slide.notes_text_frame.text.strip()
                    if notes:
                        parts.append(f"Notatki prowadzącego: {notes}")
                pages.append(Page(f.name, i, _clean("\n".join(parts))))
        elif ext == ".docx":
            import docx

            text = _clean("\n".join(p.text for p in docx.Document(str(f)).paragraphs))
            pages += [Page(f.name, i, t) for i, t in enumerate(_split(text), 1)]
        elif ext in (".txt", ".md"):
            text = _clean(f.read_text(encoding="utf-8", errors="replace"))
            pages += [Page(f.name, i, t) for i, t in enumerate(_split(text), 1)]
        elif ext in (".png", ".jpg", ".jpeg", ".webp"):
            pages.append(Page(f.name, 1, "", str(f)))
        else:
            log(f"! pomijam nieobsługiwany plik: {f.name}")
    return pages


def describe_pages(qwen: Qwen, pages: list[Page]) -> int:
    todo = [p for p in pages if p.image and not p.description]
    for p in tqdm(todo, desc="           strony-obrazy", unit="str", leave=False):
        p.description = qwen.chat(SLIDE_PROMPT, image=Path(p.image), max_tokens=1536, max_rounds=1)
    return len(todo)


def _words(text: str | None) -> set[str]:
    return set(re.findall(r"[^\W\d_]{3,}", (text or "").lower()))


def match_materials(slides: list[Slide], pages: list[Page]) -> int:
    """Łączy slajdy z nagrania ze stronami prezentacji po wspólnych słowach (opis slajdu vs tekst strony).
    Dopasowany slajd dostaje czysty tekst z prezentacji – dokładniejszy niż odczyt z wideo."""
    cands = [(p, _words(f"{p.text} {p.description or ''}")) for p in pages
             if Path(p.doc).suffix.lower() in (".pdf", ".pptx", ".png", ".jpg", ".jpeg", ".webp")]
    n = 0
    for s in slides:
        sw = _words(s.description)
        if len(sw) < 4:
            continue
        best, score = None, 0.0
        for p, pw in cands:
            inter = len(sw & pw)
            if inter >= 4 and len(pw) >= 3:
                sc = inter / min(len(sw), len(pw))
                if sc > score:
                    best, score = p, sc
        if best and score >= 0.45:
            s.material, s.material_text = best.ref, best.text or best.description
            best.matched = True
            n += 1
    return n


def material_chunks(pages: list[Page], max_chars: int) -> list[str]:
    """Strony niedopasowane do żadnego slajdu z nagrania -> osobne fragmenty dla modelu."""
    blocks = []
    for p in pages:
        if p.matched or not (p.text or p.description):
            continue
        b = f"[{p.ref}]\n{p.text}"
        if p.description:
            b += f"\n(zawartość graficzna) {p.description}"
        blocks.append(b)
    chunks, cur = [], ""
    for b in blocks:
        if len(cur) + len(b) > max_chars and cur:
            chunks.append(cur)
            cur = ""
        cur += b + "\n\n"
    return chunks + ([cur] if cur.strip() else [])


def find_materials(video: Path, extra: list[Path]) -> list[Path]:
    files = [f.resolve() for f in extra]
    for ext in AUTO_MATERIAL_EXTS:
        for cand in (video.with_suffix(ext), video.with_suffix(ext.upper())):
            if cand.exists() and cand.resolve() not in files:
                files.append(cand.resolve())
    return files


# ============================================================================
# 5. Wiedza z internetu (--web): Qwen wybiera zapytania, DuckDuckGo zwraca źródła
# ============================================================================

QUERY_PROMPT = """Na podstawie poniższych notatek z wykładu wypisz od 3 do {n} zapytań do wyszukiwarki internetowej,
które pomogą wyjaśnić, uzupełnić lub zweryfikować najważniejsze pojęcia z wykładu.
Zapytania po polsku (lub po angielsku, jeśli termin jest angielski). Jedno zapytanie na linię, bez numeracji i komentarzy.

--- NOTATKI ---
{body}"""

WEB_HEAD = """
--- WYNIKI WYSZUKIWANIA W INTERNECIE (pomocniczo) ---
{web}

Wyników użyj tylko do doprecyzowania definicji pojęć – treść wykładu jest najważniejsza.
Definicje uzupełnione z internetu oznacz dopiskiem „(źródło: internet)”.
"""

WEB_SECTION_PROMPT = """Poniżej wyniki wyszukiwania w internecie dotyczące pojęć z wykładu (krótko o wykładzie: {topic}).
Wybierz 3–8 najbardziej przydatnych źródeł do nauki i dla każdego napisz jedną linię w formacie:
- [tytuł](url) – czego można się z niego dowiedzieć w kontekście tego wykładu
Pomiń źródła niezwiązane z tematem. Bez nagłówka i bez wstępu. Po polsku.

--- WYNIKI ---
{web}"""


def web_context(qwen: Qwen, partial: list[str], max_queries: int, per_query: int) -> tuple[str, int]:
    try:
        from ddgs import DDGS
    except ImportError as e:
        raise RuntimeError("brak pakietu ddgs (pip install ddgs)") from e

    body = "\n\n".join(partial)[:MAX_CHARS]
    raw = qwen.chat(QUERY_PROMPT.format(n=max_queries, body=body), max_tokens=256, max_rounds=0)
    queries = [re.sub(r"^[\s\-*\d.)]+", "", q).strip().strip('"') for q in raw.splitlines()]
    queries = [q for q in queries if len(q) > 2][:max_queries]

    seen, blocks = set(), []
    ddg = DDGS()
    for q in tqdm(queries, desc="           szukanie", unit="zap", leave=False):
        try:
            results = ddg.text(q, region="pl-pl", max_results=per_query)
        except Exception:
            continue
        for r in results:
            url = r.get("href") or r.get("url")
            if not url or url in seen:
                continue
            seen.add(url)
            blocks.append(f"- [{r.get('title', '').strip()}]({url})\n  {r.get('body', '').strip()}")
    return "\n".join(blocks), len(queries)


# ============================================================================
# 5b. Wstawianie slajdów w treść notatki (tam, gdzie model się do nich odwołuje)
# ============================================================================

_SLIDE_REF = re.compile(
    r"\bslajd\w*\s+(?:nr\.?\s*)?(\d+(?:\s*(?:[,–—-]|\bi\b|\boraz\b|\bdo\b)\s*\d+)*)", re.IGNORECASE)
_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
_TS = re.compile(r"\[(\d{1,2}):(\d{2})(?::(\d{2}))?\]")
# sekcje, w których odwołanie do slajdu nie jest dobrym miejscem na obrazek
_SKIP_SECTIONS = re.compile(r"tl;?dr|pytania|pojęcia|dodatkowe materiały|nieomówione|źródła", re.IGNORECASE)


def _slide_numbers(match: str) -> list[int]:
    nums: list[int] = []
    for a, sep, b in re.findall(r"(\d+)\s*(?:([–—-]|\bdo\b)\s*(\d+))?", match):
        if sep and b and 0 <= int(b) - int(a) <= 30:
            nums += range(int(a), int(b) + 1)
        else:
            nums.append(int(a))
            if b:
                nums.append(int(b))
    return nums


MAX_SLIDES_PER_BLOCK = 3  # więcej obrazków pod jednym akapitem = "wysyp"; reszta idzie do sekcji na końcu


def _slide_description(sl: Slide) -> list[str]:
    """Opis slajdu (z Qwen3-VL) jako zwykły tekst notatki – nagłówki z opisu zamieniamy na pogrubienie,
    żeby nie rozbijały struktury dokumentu."""
    text = (sl.description or sl.material_text or "").strip()
    if not text:
        return []
    out = []
    for line in text.splitlines():
        m = _HEADING.match(line)
        out.append(f"**{m.group(2).strip()}**" if m else line)
    return ["", *out, ""]


def _figure(sl: Slide) -> list[str]:
    ts = fmt_ts(sl.time, False)
    src = f"slajdy/{sl.file}"
    return ["", f'<figure><img src="{src}" alt="Slajd {sl.index}">'
                f"<figcaption>Slajd {sl.index} · {ts}</figcaption></figure>", ""]


def embed_slides(notes: str, slides: list[Slide]) -> tuple[str, dict[str, int]]:
    """Wstawia obrazy slajdów w notatkę:
      1) pod blokiem (akapit/lista), w którym model pierwszy raz pisze „(slajd N)” – w części merytorycznej,
         maks. MAX_SLIDES_PER_BLOCK pod jednym blokiem,
      2) pozostałe – w sekcji „Slajdy bez odwołań w notatce” na końcu, każdy z opisem wygenerowanym przez model."""
    by_idx = {sl.index: sl for sl in slides}
    lines = notes.splitlines()
    n = len(lines)

    # podział na linie: czy w kodzie/wzorze, czy nagłówek; bloki = ciągi linii między pustymi
    # liniami i nagłówkami (nagłówek to osobny blok – modele często nie dają pustych linii)
    in_fence = [False] * n
    fence = None
    for ln, text in enumerate(lines):
        st = text.strip()
        if fence:
            in_fence[ln] = True
            if (fence == "```" and st.startswith("```")) or (fence == "$$" and st.endswith("$$")):
                fence = None
        elif st.startswith("```"):
            in_fence[ln], fence = True, "```"
        elif st.startswith("$$"):
            in_fence[ln] = True
            if not (len(st) > 4 and st.endswith("$$")):
                fence = "$$"
    is_heading = [bool(_HEADING.match(t)) and not in_fence[i] for i, t in enumerate(lines)]

    blocks: list[tuple[int, int]] = []
    start = None
    for ln in range(n):
        blank = not lines[ln].strip() and not in_fence[ln]
        if blank or is_heading[ln]:
            if start is not None:
                blocks.append((start, ln - 1))
                start = None
            if is_heading[ln]:
                blocks.append((ln, ln))
            continue
        if start is None:
            start = ln
    if start is not None:
        blocks.append((start, n - 1))

    # nagłówki: poziom, czas (jeśli jest), czy to sekcja "pomijana"
    headings: list[tuple[int, int, float | None]] = []  # (linia, poziom, czas)
    skip_line = [False] * n
    cur_skip = False
    for ln, text in enumerate(lines):
        m = _HEADING.match(text) if is_heading[ln] else None
        if m:
            level, title = len(m.group(1)), m.group(2)
            t = _TS.search(title)
            sec = None
            if t:
                h, mi, se = (int(t.group(1)), int(t.group(2)), int(t.group(3) or 0)) if t.group(3) \
                    else (0, int(t.group(1)), int(t.group(2)))
                sec = h * 3600 + mi * 60 + se
            headings.append((ln, level, sec))
            if level <= 2:  # tylko sekcje główne decydują (TL;DR, pojęcia, pytania…); ### dziedziczą
                cur_skip = bool(_SKIP_SECTIONS.search(title))
        skip_line[ln] = cur_skip

    insert_after: dict[int, list[Slide]] = {}
    placed: set[int] = set()
    stats = {"przy_odwolaniu": 0, "na_koncu": 0}

    # 1) przy pierwszym odwołaniu (pomijamy nagłówki, kod, wzory i sekcje typu TL;DR/pytania)
    for a, b in blocks:
        if is_heading[a]:
            continue
        text = "\n".join(lines[ln] for ln in range(a, b + 1)
                         if not skip_line[ln] and not in_fence[ln] and "<figure" not in lines[ln])
        for m in _SLIDE_REF.finditer(text):
            nums = _slide_numbers(m.group(1))
            if len(nums) > MAX_SLIDES_PER_BLOCK:
                continue  # szeroki zakres ("slajdy 5–12") – nie wstawiamy części, całość trafi na koniec
            for k in nums:
                if k in by_idx and k not in placed and len(insert_after.get(b, [])) < MAX_SLIDES_PER_BLOCK:
                    placed.add(k)
                    insert_after.setdefault(b, []).append(by_idx[k])
                    stats["przy_odwolaniu"] += 1

    out: list[str] = []
    for ln, text in enumerate(lines):
        out.append(text)
        for sl in sorted(insert_after.get(ln, []), key=lambda x: x.index):
            out += _figure(sl)
    rest = [sl for sl in slides if sl.index not in placed]
    if rest:
        out += ["", "---", "", "## Slajdy bez odwołań w notatce", "",
                "Slajdy pokazane na wykładzie, do których notatka się nie odwołuje – z opisem treści.", ""]
        for sl in rest:
            out += ["", f"### Slajd {sl.index} · {fmt_ts(sl.time, False)}"]
            out += _figure(sl)
            out += _slide_description(sl)
        stats["na_koncu"] = len(rest)
    return re.sub(r"\n{3,}", "\n\n", "\n".join(out)).strip() + "\n", stats


# ============================================================================
# 6. Eksport notatki do PDF (Markdown -> HTML + KaTeX -> Chromium -> PDF)
# ============================================================================

KATEX = "https://cdn.jsdelivr.net/npm/katex@0.16.11/dist"

PDF_CSS = """
@page { size: A4; margin: 18mm 16mm 20mm 16mm; }
:root { --ink:#1d1f23; --muted:#5b616b; --line:#e3e6ea; --accent:#2457c5; --code:#f5f7fa; }
html { -webkit-print-color-adjust: exact; print-color-adjust: exact; }
body { font-family: "Inter","Segoe UI","Noto Sans","DejaVu Sans",Arial,sans-serif; color: var(--ink);
       font-size: 10.5pt; line-height: 1.55; margin: 0; }
h1 { font-size: 22pt; line-height: 1.2; margin: 0 0 6mm; padding-bottom: 3mm; border-bottom: 2px solid var(--accent); }
h2 { font-size: 15pt; margin: 9mm 0 3mm; color: var(--accent); break-after: avoid; }
h3 { font-size: 12pt; margin: 6mm 0 2mm; break-after: avoid; }
h4 { font-size: 10.5pt; margin: 4mm 0 1mm; break-after: avoid; }
p, ul, ol { margin: 0 0 2.5mm; }
li { margin: 0.6mm 0; }
li > ul, li > ol { margin-top: 0.6mm; }
strong { font-weight: 650; }
a { color: var(--accent); text-decoration: none; word-break: break-word; }
hr { border: 0; border-top: 1px solid var(--line); margin: 7mm 0; }
blockquote { margin: 3mm 0; padding: 2mm 4mm; border-left: 3px solid var(--accent); background: #f3f6fc; color: var(--muted); }
code { font-family: "JetBrains Mono","Fira Code","DejaVu Sans Mono",Consolas,monospace; font-size: 9pt;
       background: var(--code); padding: 0.3mm 1mm; border-radius: 2px; }
pre { background: var(--code); border: 1px solid var(--line); border-radius: 4px; padding: 3mm 4mm;
      white-space: pre-wrap; word-break: break-word; break-inside: avoid; }
pre code { background: none; padding: 0; }
table { border-collapse: collapse; width: 100%; margin: 3mm 0; font-size: 9.5pt; break-inside: avoid; }
th, td { border: 1px solid var(--line); padding: 1.5mm 2.5mm; text-align: left; vertical-align: top; }
th { background: #f0f3f8; }
figure { margin: 3mm 0 5mm; break-inside: avoid; text-align: center; }
figure img { max-height: 80mm; margin: 0 auto 1.5mm; }
figcaption { font-size: 8.5pt; color: var(--muted); }
img { max-width: 100%; max-height: 120mm; display: block; margin: 2mm auto 5mm; border: 1px solid var(--line);
      border-radius: 3px; break-inside: avoid; }
.katex-display { margin: 3mm 0; overflow-x: auto; overflow-y: hidden; }
"""

PDF_HTML = """<!doctype html>
<html lang="pl"><head><meta charset="utf-8"><title>{title}</title>
<link rel="stylesheet" href="{katex}/katex.min.css">
<style>{css}</style>
<script defer src="{katex}/katex.min.js"></script>
<script defer src="{katex}/contrib/auto-render.min.js"
  onload="renderMathInElement(document.body,{{delimiters:[{{left:'\\\\[',right:'\\\\]',display:true}},{{left:'\\\\(',right:'\\\\)',display:false}}],throwOnError:false}});document.body.dataset.math='ok'"></script>
</head><body>
{body}
</body></html>"""


def _fix_markdown(md: str) -> str:
    """Modele często piszą listę albo wzór $$…$$ tuż pod akapitem, bez pustej linii –
    Markdown nie uznałby ich wtedy za listę / blok wzoru. Dokładamy brakujące puste linie."""
    # 1) bloki wzorów $$ … $$ (także jednoliniowe) zawsze otoczone pustymi liniami
    lines, fixed, in_math = md.splitlines(), [], False
    for line in lines:
        st = line.strip()
        opens = st.startswith("$$") and not in_math
        single = opens and len(st) > 4 and st.endswith("$$")
        if opens:
            if fixed and fixed[-1].strip():
                fixed.append("")
            fixed.append(line)
            if single:
                fixed.append("")
            else:
                in_math = True
            continue
        fixed.append(line)
        if in_math and st.endswith("$$"):
            in_math = False
            fixed.append("")
    # 2) listy tuż pod akapitem
    out, prev = [], ""
    for line in fixed:
        is_item = re.match(r"\s*([-*+]|\d+[.)])\s+", line)
        prev_item = re.match(r"\s*([-*+]|\d+[.)])\s+", prev)
        if is_item and prev.strip() and not prev_item and not prev.lstrip().startswith(("#", "|", ">")) \
                and not line.startswith((" ", "\t")):
            out.append("")
        out.append(line)
        prev = line
    return "\n".join(out)


def markdown_to_html(md: str, title: str) -> str:
    import html

    import markdown

    body = markdown.markdown(
        _fix_markdown(md),
        extensions=["extra", "sane_lists", "pymdownx.arithmatex", "pymdownx.tilde"],
        extension_configs={"pymdownx.arithmatex": {"generic": True}},
    )
    return PDF_HTML.format(title=html.escape(title), css=PDF_CSS, katex=KATEX, body=body)


def _system_browser() -> str | None:
    import shutil

    for name in ("chromium", "chromium-browser", "google-chrome", "google-chrome-stable", "chrome",
                 "microsoft-edge", "microsoft-edge-stable", "brave", "brave-browser"):
        if p := shutil.which(name):
            return p
    for p in ("/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
              "/Applications/Chromium.app/Contents/MacOS/Chromium",
              "/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge",
              r"C:\Program Files\Google\Chrome\Application\chrome.exe",
              r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
              r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"):
        if os.path.exists(p):
            return p
    return None


def _launch_browser(pw):
    """Najpierw przeglądarka z systemu (bez pobierania), potem Chromium Playwrighta w ./models/playwright."""
    import subprocess

    exe = _system_browser()
    if exe:
        try:
            return pw.chromium.launch(executable_path=exe), Path(exe).name
        except Exception as e:
            log(f"! {Path(exe).name} nie wystartował ({err_msg(e)}) – używam Chromium z Playwrighta")
    try:
        return pw.chromium.launch(), "chromium (playwright)"
    except Exception:
        pass
    log("Pobieram Chromium do eksportu PDF → models/playwright (jednorazowo)…")
    env = {**os.environ, "PLAYWRIGHT_BROWSERS_PATH": str(MODELS_DIR / "playwright")}
    cmd = [sys.executable, "-m", "playwright", "install", "--only-shell", "chromium"]
    if subprocess.run(cmd, env=env).returncode != 0 and \
            subprocess.run(cmd[:-2] + ["chromium"], env=env).returncode != 0:
        raise RuntimeError("nie mogę pobrać Chromium (brak internetu?). Zainstaluj Chrome/Chromium "
                           "albo uruchom: python -m playwright install chromium")
    return pw.chromium.launch(), "chromium (playwright)"


def markdown_to_pdf(md_path: Path, pdf_path: Path | None = None) -> tuple[Path, str]:
    """Notatka .md -> ładny PDF (wzory LaTeX, tabele, obrazki slajdów, numery stron)."""
    os.environ.setdefault("PLAYWRIGHT_BROWSERS_PATH", str(MODELS_DIR / "playwright"))
    from playwright.sync_api import sync_playwright

    md_path = md_path.resolve()
    pdf_path = (pdf_path or md_path.with_suffix(".pdf")).resolve()
    md = md_path.read_text(encoding="utf-8")
    m = re.search(r"^#\s+(.+)$", md, re.MULTILINE)
    title = m.group(1).strip() if m else md_path.stem
    # HTML obok .md, żeby względne ścieżki do slajdów (slajdy/…jpg) działały
    html_path = md_path.with_name(f".{md_path.stem}.pdf.html")
    html_path.write_text(markdown_to_html(md, title), encoding="utf-8")
    try:
        with sync_playwright() as pw:
            browser, used = _launch_browser(pw)
            try:
                page = browser.new_page()
                try:
                    page.goto(html_path.as_uri(), wait_until="networkidle", timeout=60000)
                except Exception:
                    page.goto(html_path.as_uri(), wait_until="load", timeout=60000)
                try:  # wzory renderuje KaTeX (z CDN); bez internetu zostaną jako tekst LaTeX
                    page.wait_for_function("document.body.dataset.math === 'ok'", timeout=15000)
                except Exception:
                    used += ", wzory jako tekst (brak KaTeX/internetu)"
                page.pdf(
                    path=str(pdf_path), format="A4", print_background=True, prefer_css_page_size=True,
                    display_header_footer=True, header_template="<div></div>",
                    footer_template='<div style="width:100%;font-size:8px;color:#8a9099;text-align:center;">'
                                    '<span class="pageNumber"></span> / <span class="totalPages"></span></div>',
                )
            finally:
                browser.close()
    finally:
        html_path.unlink(missing_ok=True)
    return pdf_path, used


# ============================================================================
# Zapis / cache
# ============================================================================

def write_transcript(path: Path, video: Path, segments: list[Segment], model: str, lang: str) -> None:
    speakers = {s.speaker for s in segments if s.speaker}
    head = [f"# Transkrypcja: {video.name}",
            f"# Model: {model} | Język: {lang} | Długość: {fmt_ts(segments[-1].end if segments else 0, False)}"
            + (f" | Mówcy: {len(speakers)}" if speakers else ""), ""]
    body = [f"[{fmt_ts(s.start)} --> {fmt_ts(s.end)}] {s.speaker + ': ' if s.speaker else ''}{s.text}"
            for s in segments]
    path.write_text("\n".join(head + body) + "\n", encoding="utf-8")


def write_slides_md(path: Path, slides: list[Slide]) -> None:
    out = ["# Slajdy\n"]
    for s in slides:
        out += [f"## Slajd {s.index} — {fmt_ts(s.time, False)}\n", f"![Slajd {s.index}](slajdy/{s.file})\n"]
        if s.description:
            out.append(s.description.strip() + "\n")
    path.write_text("\n".join(out), encoding="utf-8")


def save_json(path: Path, items) -> None:
    path.write_text(json.dumps([asdict(i) for i in items], ensure_ascii=False, indent=1), "utf-8")


def load_json(path: Path, cls):
    """Wczytuje cache; pola nieznane tej wersji skryptu (np. z innej wersji) są pomijane."""
    from dataclasses import fields

    known = {f.name for f in fields(cls)}
    return [cls(**{k: v for k, v in d.items() if k in known}) for d in json.loads(path.read_text("utf-8"))]


# ============================================================================
# Pipeline + CLI
# ============================================================================

def process(video: Path, args, device: str) -> None:
    t0 = time.time()
    out = video.parent / f"{video.stem}_wyniki"
    out.mkdir(exist_ok=True)
    seg_cache, slides_cache, slides_dir = out / ".segmenty.json", out / ".slajdy.json", out / "slajdy"
    is_audio = video.suffix.lower() in AUDIO_EXTS

    segments: list[Segment] = load_json(seg_cache, Segment) if seg_cache.exists() and not args.force else []
    has_speakers = any(s.speaker for s in segments)
    # cache wystarcza, chyba że teraz chcemy mówców, a w cache ich nie ma
    cached = bool(segments) and (has_speakers or not args.speakers)

    mat_files = find_materials(video, args.materials) if args.summary else []
    total = (3 + (1 if args.speakers else 0) + (1 if args.extract else 0)
             + (5 if args.summary else 0) + (1 if args.summary and not args.no_pdf else 0)
             + (2 if mat_files else 0) + (1 if args.web else 0))
    st = Stages(total)
    audio = None

    # ---- transkrypcja ----
    with st.stage("Dekodowanie audio") as s:
        if cached:
            s.skip("transkrypcja z cache")
        else:
            audio = load_audio(video)
            s.note = f"{fmt_dur(len(audio) / SAMPLE_RATE)} nagrania"

    with st.stage(f"Transkrypcja (Whisper {args.model}, {device})") as s:
        if segments:
            s.skip(f"z cache, {len(segments)} segmentów")
        else:
            segments = transcribe(audio, args.model, args.language, device, args.prompt)
            s.note = f"{len(segments)} segmentów"

    if args.speakers:
        with st.stage("Rozpoznawanie mówców") as s:
            if cached:
                s.skip(f"z cache, mówców: {len({x.speaker for x in segments})}")
            else:
                try:
                    turns = diarize(audio, device, args.hf_token,
                                    args.num_speakers, args.min_speakers, args.max_speakers)
                    segments = assign_speakers(segments, turns)
                    s.note = f"mówców: {len({x.speaker for x in segments})}"
                except Exception as e:
                    s.skip(str(e))

    if audio is not None:
        del audio
        free_memory()
        save_json(seg_cache, segments)

    with st.stage("Zapis transkrypcji") as s:
        txt = out / f"{video.stem}_transkrypcja.txt"
        write_transcript(txt, video, segments, args.model, args.language)
        s.note = str(txt.relative_to(video.parent))

    # ---- slajdy ----
    slides: list[Slide] = []
    if args.extract:
        with st.stage("Wydobywanie klatek slajdów") as s:
            if is_audio:
                s.skip("plik audio – brak obrazu")
            elif slides_cache.exists() and not args.force:
                slides = load_json(slides_cache, Slide)
                s.skip(f"z cache, {len(slides)} slajdów")
            else:
                slides = extract_slides(video, slides_dir, args.slide_interval, args.slide_threshold)
                save_json(slides_cache, slides)
                s.note = f"{len(slides)} unikalnych slajdów → slajdy/"

    # ---- materiały (prezentacja, notatki) ----
    pages: list[Page] = []
    if mat_files:
        with st.stage("Wczytywanie materiałów") as s:
            pages = load_materials(mat_files, out / ".materialy")
            s.note = f"{len(mat_files)} plików ({', '.join(f.name for f in mat_files)}), {len(pages)} stron/części"

    # ---- Qwen3-VL ----
    if args.summary:
        qwen = None
        try:
            with st.stage("Ładowanie modelu Qwen3-VL") as s:
                qwen = Qwen(args.qwen, device)
                s.note = f"{qwen.name} ({qwen.backend})"

            with st.stage("Opisywanie slajdów") as s:
                todo = [x for x in slides if not x.description]
                if not slides:
                    s.skip("brak slajdów")
                elif not todo:
                    s.skip("opisy z cache")
                else:
                    describe_slides(qwen, slides, slides_dir)
                    save_json(slides_cache, slides)
                    s.note = f"{len(todo)} opisów"

            if mat_files:
                with st.stage("Dopasowanie materiałów do slajdów z nagrania") as s:
                    described = describe_pages(qwen, pages)
                    for x in slides:  # dopasowanie liczone od nowa przy każdym uruchomieniu
                        x.material = x.material_text = None
                    matched = match_materials(slides, pages)
                    s.note = (f"{matched}/{len(slides)} slajdów dopasowanych, "
                              f"{sum(not p.matched for p in pages)} stron jako materiał dodatkowy"
                              + (f", {described} stron-obrazów opisanych" if described else ""))

            with st.stage("Przekazywanie transkrypcji, slajdów i materiałów do modelu"
                          if mat_files else "Przekazywanie transkrypcji i slajdów do modelu") as s:
                partial, reused = partial_notes(qwen, segments, slides, pages,
                                                cache=None if args.force else out / ".notatki_czesciowe.json")
                s.note = f"{len(partial)} fragmentów przeanalizowanych" + (f" ({reused} z cache)" if reused else "")

            web = ""
            if args.web:
                with st.stage("Wyszukiwanie w internecie") as s:
                    try:
                        web, nq = web_context(qwen, partial, args.web_queries, args.web_results)
                        if web:
                            s.note = f"{nq} zapytań, {web.count('- [')} źródeł"
                        else:
                            s.skip("brak wyników (brak internetu albo limit DuckDuckGo)")
                    except Exception as e:
                        s.skip(f"pominięto: {e}")

            with st.stage("Generowanie podsumowania") as s:
                notes = final_notes(qwen, partial, web)
                s.note = f"{len(notes.split())} słów"
        finally:
            if qwen:
                qwen.close()

        with st.stage("Zapis notatki") as s:
            placement = ""
            if slides:
                notes, st_ = embed_slides(notes, slides)
                placement = f" — slajdy: {st_['przy_odwolaniu']} w treści, {st_['na_koncu']} w sekcji na końcu"
            note = out / f"{video.stem}_notatka.md"
            note.write_text(notes, encoding="utf-8")
            s.note = str(note.relative_to(video.parent)) + placement

        if not args.no_pdf:
            with st.stage("Eksport notatki do PDF") as s:
                try:
                    pdf, used = markdown_to_pdf(note)
                    s.note = f"{pdf.relative_to(video.parent)} ({used})"
                except Exception as e:
                    s.skip(f"nie udało się: {err_msg(e)} – notatka .md jest gotowa")

    if slides:
        write_slides_md(out / "slajdy.md", slides)
    print(_c("1;32", f"\n✓ {video.name} gotowe w {fmt_dur(time.time() - t0)}") + f"  →  {out}")


def parse_args(argv=None) -> argparse.Namespace:
    p = argparse.ArgumentParser(
        prog="transcriber.py",
        description="Lokalna transkrypcja nagrań (PL) z mówcami, slajdami i notatką do nauki. "
                    "Modele pobierają się przy pierwszym użyciu do ./models.",
    )
    p.add_argument("paths", nargs="+", type=Path,
                   help="pliki wideo/audio lub foldery z nimi; plik .md = tylko konwersja gotowej notatki do PDF")
    p.add_argument("--speakers", action="store_true",
                   help="podział na mówców (pyannote community-1, bez tokenu)")
    p.add_argument("--extract", action="store_true", help="zapisz slajdy (istotne zmiany obrazu) jako JPG")
    p.add_argument("--summary", action="store_true", help="opisz slajdy i zrób notatkę .md (włącza --extract)")
    p.add_argument("--materials", nargs="+", type=Path, default=[], metavar="PLIK",
                   help="materiały do notatki: prezentacja (.pdf/.pptx), notatki (.docx/.txt/.md), obrazy; "
                        "plik .pdf/.pptx o tej samej nazwie co nagranie jest dołączany automatycznie; włącza --summary")
    p.add_argument("--all", action="store_true",
                   help="wszystko naraz: --extract --summary --web --speakers")
    p.add_argument("--web", action="store_true",
                   help="uzupełnij notatkę wiedzą z internetu (DuckDuckGo, bez klucza); włącza --summary")

    g = p.add_argument_group("transkrypcja")
    g.add_argument("--model", default=DEFAULT_WHISPER,
                   help="Whisper: large-v3-turbo (domyślny, szybki) | large-v3 (najlepszy PL) | medium | small")
    g.add_argument("--language", default="pl")
    g.add_argument("--prompt", help="temat/terminologia wykładu – mniej błędów w nazwach")
    g.add_argument("--device", default="auto", choices=["auto", "cuda", "mps", "cpu"])

    g = p.add_argument_group("mówcy (z --speakers)")
    g.add_argument("--num-speakers", type=int, help="dokładna liczba mówców, jeśli znana")
    g.add_argument("--min-speakers", type=int)
    g.add_argument("--max-speakers", type=int)
    g.add_argument("--hf-token", help="opcjonalny token HF – tylko gdy publiczna kopia modelu jest niedostępna")

    g = p.add_argument_group("slajdy / notatka")
    g.add_argument("--slide-threshold", type=float, default=0.35,
                   help="jaka część treści slajdu musi być nowa, by zapisać nowy slajd "
                        "(0.35 = 35%%; mniej = łapie też dopisywane punkty)")
    g.add_argument("--slide-interval", type=float, default=1.0, help="co ile sekund sprawdzać obraz")
    g.add_argument("--qwen", default="auto",
                   help="rozmiar Qwen3-VL: auto | 2b | 4b | 8b (albo repozytorium HF)")
    g.add_argument("--web-queries", type=int, default=6, help="maks. liczba zapytań do wyszukiwarki")
    g.add_argument("--web-results", type=int, default=3, help="wyników na zapytanie")

    p.add_argument("--no-pdf", action="store_true", help="nie twórz PDF z notatki (domyślnie powstaje z --summary)")
    p.add_argument("--force", action="store_true", help="licz wszystko od nowa (ignoruj cache)")
    a = p.parse_args(argv)
    if a.all:
        a.extract = a.summary = a.web = a.speakers = True
    for m in a.materials:
        if not m.is_file():
            p.error(f"nie ma pliku materiałów: {m}")
    a.summary = a.summary or a.web or bool(a.materials)
    a.extract = a.extract or a.summary
    return a


def collect_files(paths: list[Path]) -> list[Path]:
    exts = VIDEO_EXTS | AUDIO_EXTS | {".md"}
    files = []
    for path in paths:
        if path.is_dir():
            files += sorted(f for f in path.iterdir() if f.suffix.lower() in exts - {".md"})
        elif path.is_file():
            files.append(path)
        else:
            print(f"! Nie istnieje: {path}")
    return files


def main(argv=None) -> int:
    args = parse_args(argv)
    files = collect_files(args.paths)
    if not files:
        print("Brak plików do przetworzenia.")
        return 1
    # tryb "tylko PDF": python transcriber.py wyklad_wyniki/wyklad_notatka.md
    md_files = [f for f in files if f.suffix.lower() == ".md"]
    files = [f for f in files if f.suffix.lower() != ".md"]
    failed = 0
    for f in md_files:
        print(_c("1", f"\n━━━ PDF: {f.name} ━━━"))
        with Stages(1).stage("Eksport notatki do PDF") as s:
            try:
                pdf, used = markdown_to_pdf(f)
                s.note = f"{pdf} ({used})"
            except Exception as e:
                failed += 1
                s.skip(f"nie udało się: {err_msg(e, 300)}")
    if not files:
        return 1 if failed else 0

    device = detect_device(args.device)
    mode = " + ".join(["transkrypcja"] + [n for f, n in [(args.speakers, "mówcy"), (args.extract, "slajdy"),
                                                         (args.summary, "notatka"), (bool(args.materials), "materiały"),
                                                         (args.web, "internet")] if f])
    print(_c("1;36", "transcriber") + f"  urządzenie: {device} | tryb: {mode} | modele: {MODELS_DIR}")
    if device == "cuda" and (args.summary or args.speakers):
        ok, why = torch_cuda_ok()
        if not ok:
            print(_c("1;33", f"! Whisper użyje karty, ale PyTorch nie: {why}.") + "\n"
                  "  Qwen i rozpoznawanie mówców pójdą na CPU (wolno). Naprawa: pip install torch torchvision "
                  "--index-url https://download.pytorch.org/whl/cu128")
    for i, f in enumerate(files, 1):
        print(_c("1", f"\n━━━ [{i}/{len(files)}] {f.name} ━━━"))
        try:
            process(f.resolve(), args, device)
        except KeyboardInterrupt:
            print("\nPrzerwano.")
            return 130
        except Exception as e:
            failed += 1
            print(_c("1;31", f"✗ {f.name}: {err_msg(e, 400)}"))
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
